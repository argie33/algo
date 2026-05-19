#!/usr/bin/env python3
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

"""
NAAIM Exposure Index Loader.

Downloads the weekly NAAIM (National Association of Active Investment Managers)
Exposure Index data and stores it in the naaim table.

NAAIM publishes their historical data as a public spreadsheet.
Released every Wednesday.

Run:
    python3 loadnaaim.py
"""

import logging
from datetime import date, datetime
from typing import List, Optional

import requests

from config.env_loader import load_env
from utils.db_connection import get_db_connection

log = logging.getLogger(__name__)

NAAIM_PAGE = "https://www.naaim.org/programs/naaim-exposure-index/"
_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}


def _discover_naaim_url() -> Optional[str]:
    """Scrape the NAAIM page to find the current spreadsheet download link."""
    import re
    try:
        resp = requests.get(NAAIM_PAGE, headers=_HEADERS, timeout=20)
        resp.raise_for_status()
        links = re.findall(r'https?://[^\s"\']*?USE_Data[^\s"\']*\.xlsx', resp.text, re.IGNORECASE)
        if links:
            return links[0]
        # Fallback: any xlsx on naaim.org
        links = re.findall(r'https?://(?:www\.)?naaim\.org[^\s"\']*\.xlsx', resp.text, re.IGNORECASE)
        return links[0] if links else None
    except Exception as e:
        log.warning("NAAIM page scrape failed: %s", e)
        return None


def fetch_naaim() -> List[dict]:
    """Download and parse NAAIM exposure index spreadsheet."""
    url = _discover_naaim_url()
    if not url:
        log.error("NAAIM: could not find download URL on NAAIM website")
        return []

    log.info("NAAIM: downloading from %s", url)
    content = None
    try:
        resp = requests.get(url, timeout=60, headers=_HEADERS)
        if resp.status_code == 200:
            content = resp.content
    except Exception as e:
        log.error("NAAIM download failed: %s", e)
        return []

    if not content:
        log.error("NAAIM: download returned empty content")
        return []

    # Try openpyxl (xlsx) first, then xlrd (xls)
    rows = _parse_xlsx(content) or _parse_xls(content)
    log.info("NAAIM: parsed %d rows", len(rows))
    return rows


def _parse_xlsx(content: bytes) -> List[dict]:
    try:
        from openpyxl import load_workbook
        import io
        wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        header_skipped = False
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            # Skip header row (Date, Mean/Average, ...)
            if not header_skipped and isinstance(row[0], str):
                header_skipped = True
                continue
            try:
                dt = _parse_date(row[0])
                if dt is None:
                    continue
                # NAAIM file columns: Date, Mean/Average, Most Bearish, Quart1, Quart2, Quart3
                mean_val = float(row[1]) if row[1] is not None else None
                bearish = float(row[2]) if len(row) > 2 and row[2] is not None else None
                # Quart3 (75th pct) as proxy for bullish sentiment
                bullish = float(row[5]) if len(row) > 5 and row[5] is not None else None
                if mean_val is None:
                    continue
                rows.append({
                    "date": str(dt),
                    "naaim_number_mean": mean_val,
                    "bullish": bullish,
                    "bearish": bearish,
                })
            except (ValueError, TypeError):
                continue
        return rows
    except Exception as e:
        log.debug("openpyxl parse failed: %s", e)
        return []


def _parse_xls(content: bytes) -> List[dict]:
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=content)
        ws = wb.sheet_by_index(0)
        rows = []
        for row_idx in range(ws.nrows):
            try:
                dt = _parse_xlrd_date(ws.cell(row_idx, 0), wb.datemode)
                if dt is None:
                    continue
                mean_val = float(ws.cell(row_idx, 1).value)
                bullish = float(ws.cell(row_idx, 2).value) if ws.ncols > 2 else None
                bearish = float(ws.cell(row_idx, 3).value) if ws.ncols > 3 else None
                rows.append({
                    "date": str(dt),
                    "naaim_number_mean": mean_val,
                    "bullish": bullish,
                    "bearish": bearish,
                })
            except (ValueError, TypeError, IndexError):
                continue
        return rows
    except Exception as e:
        log.debug("xlrd parse failed: %s", e)
        return []


def _parse_date(val) -> Optional[date]:
    if isinstance(val, date):
        return val
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, str):
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y%m%d"):
            try:
                return datetime.strptime(val.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _parse_xlrd_date(cell, datemode: int) -> Optional[date]:
    import xlrd
    if cell.ctype == xlrd.XL_CELL_DATE:
        t = xlrd.xldate_as_tuple(cell.value, datemode)
        return date(*t[:3])
    if cell.ctype in (xlrd.XL_CELL_TEXT,):
        return _parse_date(cell.value)
    return None


def upsert_rows(rows: List[dict]) -> int:
    if not rows:
        return 0
    conn = get_db_connection()
    cur = conn.cursor()
    inserted = 0
    try:
        for row in rows:
            cur.execute(
                """
                INSERT INTO naaim (date, naaim_number_mean, bullish, bearish)
                VALUES (%(date)s, %(naaim_number_mean)s, %(bullish)s, %(bearish)s)
                ON CONFLICT (date) DO UPDATE SET
                    naaim_number_mean = EXCLUDED.naaim_number_mean,
                    bullish = EXCLUDED.bullish,
                    bearish = EXCLUDED.bearish
                """,
                row,
            )
            inserted += cur.rowcount
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()
    return inserted


def main() -> int:
    load_env()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    rows = fetch_naaim()
    if not rows:
        log.warning("No NAAIM data fetched — source may be unavailable")
        return 0  # Non-fatal: NAAIM site can be unavailable

    inserted = upsert_rows(rows)
    log.info("NAAIM: %d rows upserted", inserted)
    return 0


if __name__ == "__main__":
    sys.exit(main())
